# -*- coding: utf-8 -*-
import re
import zipfile
from xml.etree import ElementTree as ET
from docx import Document

METHOD_PATH = r"d:\Downloads\Telegram Desktop\ОФОРМЛЕНИЕ КП и ДП 2026.docx"
PZ_PATH = r"c:\Users\user\Desktop\PerfumeStore_NEW\PerfumeStoreCourseProject\Documentation\ПЗКурсовая.docx"
OUT_PATH = r"c:\Users\user\Desktop\PerfumeStore_NEW\PerfumeStoreCourseProject\Documentation\pz_deep_analysis.txt"

NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def extract_sections():
    doc = Document(METHOD_PATH)
    sections = {
        "2.4": False,
        "2.5": False,
        "2.6": False,
        "2.7": False,
    }
    current = None
    lines = []
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        for key in sections:
            if re.match(rf"^{re.escape(key)}\b", t):
                current = key
        if current and t:
            if current == "2.4" and re.match(r"^2\.5\b", t):
                current = None
            elif current == "2.5" and re.match(r"^2\.6\b", t):
                current = None
            elif current == "2.6" and re.match(r"^2\.7\b", t):
                current = None
            elif current == "2.7" and re.match(r"^2\.8\b", t):
                current = None
            if current:
                lines.append(f"[{current}] {t}")
    return lines


def analyze_pz():
    doc = Document(PZ_PATH)
    lines = []
    lines.append(f"Total paragraphs: {len(doc.paragraphs)}")
    lines.append(f"Total tables: {len(doc.tables)}")

    # styles distribution
    from collections import Counter
    styles = Counter(p.style.name for p in doc.paragraphs if p.text.strip())
    lines.append("Style distribution:")
    for s, c in styles.most_common(20):
        lines.append(f"  {s}: {c}")

    # special sections
    keywords = [
        "СОДЕРЖАНИЕ", "ВВЕДЕНИЕ", "ЗАКЛЮЧЕНИЕ", "СПИСОК",
        "ПРИЛОЖЕНИЕ", "Рисунок", "Таблица", "РЕФЕРАТ",
    ]
    lines.append("\nSpecial paragraphs:")
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if any(k in t for k in keywords) and (len(t) < 120 or t.startswith("Рисунок") or t.startswith("Таблица")):
            pf = p.paragraph_format
            r = p.runs[0] if p.runs else None
            lines.append(
                f"{i:4d}|{p.style.name}|align={p.alignment}|sp_b={pf.space_before}|sp_a={pf.space_after}|ind={pf.first_line_indent}|bold={r.font.bold if r else None}| {t[:100]}"
            )

    # line spacing sample
    lines.append("\nLine spacing samples (body):")
    count = 0
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if len(t) > 80 and p.style.name == "Normal":
            pf = p.paragraph_format
            lines.append(
                f"{i}|rule={pf.line_spacing_rule}|spacing={pf.line_spacing}|font={p.runs[0].font.name if p.runs else None}|size={p.runs[0].font.size if p.runs else None}"
            )
            count += 1
            if count >= 15:
                break

    # bullets and lists
    lines.append("\nList/bullet paragraphs:")
    bullet_count = 0
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if not t:
            continue
        if t.startswith(("•", "−", "–", "-", "", "\uf0b7")) or "List" in (p.style.name or ""):
            lines.append(f"{i:4d}|{p.style.name}| {t[:90]}")
            bullet_count += 1
            if bullet_count >= 40:
                break

    # headings spacing
    lines.append("\nHeadings with spacing:")
    for i, p in enumerate(doc.paragraphs):
        if p.style.name in ("Заголовок ПЗ 1", "Заголовок ПЗ 2", "Heading 1", "Heading 2"):
            pf = p.paragraph_format
            lines.append(
                f"{i:4d}|{p.style.name}|sp_b={pf.space_before}|sp_a={pf.space_after}|ind={pf.first_line_indent}|align={p.alignment}| {p.text.strip()[:80]}"
            )

    # figure/table count
    figs = sum(1 for p in doc.paragraphs if p.text.strip().startswith("Рисунок"))
    tabs = sum(1 for p in doc.paragraphs if p.text.strip().startswith("Таблица"))
    lines.append(f"\nFigure captions: {figs}, Table captions: {tabs}")

    # check page numbers in footer
    with zipfile.ZipFile(PZ_PATH) as z:
        if "word/footer1.xml" in z.namelist():
            footer = z.read("word/footer1.xml").decode("utf-8", errors="replace")
            lines.append("\nFooter snippet:")
            lines.append(footer[:1500])
        else:
            lines.append("\nNo footer1.xml")

    return lines


def main():
    out = []
    out.append("=== METHODOLOGY SECTIONS 2.4-2.7 ===")
    out.extend(extract_sections())
    out.append("\n=== PZ ANALYSIS ===")
    out.extend(analyze_pz())
    with open(OUT_PATH, "w", encoding="utf-8") as f:
        f.write("\n".join(out))
    print("Done:", OUT_PATH)


if __name__ == "__main__":
    main()
