# -*- coding: utf-8 -*-
import re
from collections import Counter
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

path = r"c:\Users\user\Desktop\PerfumeStore_NEW\PerfumeStoreCourseProject\Documentation\ПЗКурсовая.docx"
out = r"c:\Users\user\Desktop\PerfumeStore_NEW\PerfumeStoreCourseProject\Documentation\validation_report.txt"
doc = Document(path)
lines = []

# margins
s = doc.sections[0]
lines.append(f"Margins L={s.left_margin} R={s.right_margin} T={s.top_margin} B={s.bottom_margin}")

issues = []
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if not t or i <= 49:
        continue
    pf = p.paragraph_format
    if t.startswith("Рисунок"):
        if p.alignment != WD_ALIGN_PARAGRAPH.CENTER:
            issues.append(f"fig align {i}: {t[:60]}")
        if " — " in t:
            issues.append(f"fig emdash {i}")
        if pf.first_line_indent and pf.first_line_indent > 0:
            issues.append(f"fig indent {i}")
    if p.style.name == "Заголовок ПЗ 1" or re.match(r"^\d", t):
        if t.lower() in ("введение", "заключение") or re.match(r"^\d", t):
            bold = any(r.bold for r in p.runs if r.text.strip())
            if not bold and i >= 59:
                issues.append(f"heading not bold {i}: {t[:50]}")
    if t.startswith("•"):
        issues.append(f"bullet {i}")
    if t.startswith("СОДЕРЖАНИЕ"):
        if p.alignment != WD_ALIGN_PARAGRAPH.CENTER:
            issues.append("content not centered")
        if not any(r.bold for r in p.runs):
            issues.append("content not bold")
    if 510 <= i <= 522 and t and not t.startswith("Список"):
        if not re.match(r"^\d+\s", t):
            issues.append(f"source not numbered {i}: {t[:40]}")

lines.append(f"Issues: {len(issues)}")
lines.extend(issues[:60])

lines.append("\nTOC area 50-58:")
for i in range(50, 59):
    lines.append(f"{i}: {repr(doc.paragraphs[i].text[:100])}")

lines.append("\nHeadings sample:")
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if p.style.name == "Заголовок ПЗ 1" and t:
        b = any(r.bold for r in p.runs)
        lines.append(f"{i}|b={b}|ind={p.paragraph_format.first_line_indent}|sa={p.paragraph_format.space_after}| {t[:70]}")

lines.append("\nSources:")
for i in range(508, 525):
    if i < len(doc.paragraphs):
        lines.append(f"{i}: {doc.paragraphs[i].text[:100]}")

# spacing check body
bad_spacing = 0
for i, p in enumerate(doc.paragraphs[60:200], 60):
    t = p.text.strip()
    if len(t) > 50 and p.style.name == "Normal":
        pf = p.paragraph_format
        if pf.line_spacing_rule is not None and str(pf.line_spacing) not in ("1.0", "None"):
            bad_spacing += 1
lines.append(f"\nNon-single spacing in sample: {bad_spacing}")

with open(out, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))
print("written", out, "issues", len(issues))
