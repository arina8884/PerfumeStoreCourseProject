# -*- coding: utf-8 -*-
import zipfile
import re
from xml.etree import ElementTree as ET
from docx import Document
from docx.shared import Pt, Cm, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING

NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}

METHOD_PATH = r"d:\Downloads\Telegram Desktop\ОФОРМЛЕНИЕ КП и ДП 2026.docx"
PZ_PATH = r"c:\Users\user\Desktop\PerfumeStore_NEW\PerfumeStoreCourseProject\Documentation\ПЗКурсовая.docx"
OUT_PATH = r"c:\Users\user\Desktop\PerfumeStore_NEW\PerfumeStoreCourseProject\Documentation\analysis_output.txt"


def twips_to_mm(val):
    return round(float(val) / 56.7, 1)


def get_section_margins(path):
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml")
    root = ET.fromstring(xml)
    results = []
    for sect in root.findall(".//w:sectPr", NS):
        pgMar = sect.find("w:pgMar", NS)
        pgSz = sect.find("w:pgSz", NS)
        if pgMar is None:
            continue
        item = {}
        for k in ("top", "bottom", "left", "right", "header", "footer", "gutter"):
            v = pgMar.get(f"{{{NS['w']}}}{k}")
            if v:
                item[k] = twips_to_mm(v)
        if pgSz is not None:
            w = pgSz.get(f"{{{NS['w']}}}w")
            h = pgSz.get(f"{{{NS['w']}}}h")
            if w:
                item["page_w_mm"] = twips_to_mm(w)
            if h:
                item["page_h_mm"] = twips_to_mm(h)
        results.append(item)
    return results


def para_info(p):
    pf = p.paragraph_format
    r = p.runs[0] if p.runs else None
    return {
        "style": p.style.name if p.style else "",
        "text": p.text[:120],
        "align": str(p.alignment),
        "line_spacing": str(pf.line_spacing),
        "line_rule": str(pf.line_spacing_rule),
        "space_before": str(pf.space_before),
        "space_after": str(pf.space_after),
        "first_indent": str(pf.first_line_indent),
        "left_indent": str(pf.left_indent),
        "font": r.font.name if r else None,
        "size": str(r.font.size) if r and r.font.size else None,
        "bold": r.font.bold if r else None,
    }


def main():
    lines = []
    for label, path in [("METHOD", METHOD_PATH), ("PZ", PZ_PATH)]:
        lines.append(f"=== {label}: {path} ===")
        lines.append(f"Margins: {get_section_margins(path)}")
        doc = Document(path)
        lines.append(f"Paragraphs: {len(doc.paragraphs)}, Tables: {len(doc.tables)}")
        lines.append("")

        if label == "METHOD":
            lines.append("--- Section 2 formatting requirements ---")
            capture = False
            for i, p in enumerate(doc.paragraphs):
                t = p.text.strip()
                if re.match(r"^2\.3\b", t) or "2.3 Изложение текста" in t:
                    capture = True
                if capture and re.match(r"^2\.4\b", t):
                    capture = False
                if capture and t:
                    lines.append(f"{i:4d}| {t}")

            lines.append("")
            lines.append("--- Section 2.1 general ---")
            capture = False
            for i, p in enumerate(doc.paragraphs):
                t = p.text.strip()
                if re.match(r"^2\.1\b", t) or "2.1 Общие положения" in t:
                    capture = True
                if capture and re.match(r"^2\.2\b", t):
                    capture = False
                if capture and t:
                    lines.append(f"{i:4d}| {t}")

            lines.append("")
            lines.append("--- Section 2.2 rubrication ---")
            capture = False
            for i, p in enumerate(doc.paragraphs):
                t = p.text.strip()
                if re.match(r"^2\.2\b", t) or "2.2 Рубрикация" in t:
                    capture = True
                if capture and re.match(r"^2\.3\b", t):
                    capture = False
                if capture and t:
                    lines.append(f"{i:4d}| {t}")

            lines.append("")
            lines.append("--- Section 5 title list ---")
            capture = False
            for i, p in enumerate(doc.paragraphs):
                t = p.text.strip()
                if re.match(r"^5\b", t) and "титульн" in t.lower():
                    capture = True
                if capture and re.match(r"^ПРИЛОЖЕНИЕ", t):
                    capture = False
                if capture and t:
                    lines.append(f"{i:4d}| {t}")

            lines.append("")
            lines.append("--- Sources / bibliography ---")
            for i, p in enumerate(doc.paragraphs):
                t = p.text.strip()
                if "источник" in t.lower() and ("список" in t.lower() or "литератур" in t.lower()):
                    lines.append(f"{i:4d}| {t}")
                    for j in range(i, min(i + 25, len(doc.paragraphs))):
                        lines.append(f"  {j:4d}| {doc.paragraphs[j].text.strip()}")

        if label == "PZ":
            lines.append("--- Document structure (headings) ---")
            for i, p in enumerate(doc.paragraphs):
                t = p.text.strip()
                if not t:
                    continue
                st = p.style.name
                if st.startswith("Heading") or re.match(r"^(ВВЕДЕНИЕ|ЗАКЛЮЧЕНИЕ|СОДЕРЖАНИЕ|СПИСОК|ПРИЛОЖЕНИЕ|РЕФЕРАТ|РАЗРАБОТ)", t):
                    lines.append(f"{i:4d}| {st:20s}| {t[:100]}")
                if re.match(r"^\d+(\.\d+)*\s", t) and len(t) < 120:
                    lines.append(f"{i:4d}| {st:20s}| {t[:100]}")

            lines.append("")
            lines.append("--- Formatting issues sample ---")
            issues = []
            for i, p in enumerate(doc.paragraphs):
                t = p.text.strip()
                if not t:
                    continue
                info = para_info(p)
                # detect bullets
                if t.startswith("•") or t.startswith("\uf0b7") or "•" in t[:3]:
                    issues.append(f"bullet@{i}: {t[:80]}")
                if info["font"] and info["font"] not in (None, "Times New Roman", "Calibri"):
                    if "Times" not in str(info["font"]):
                        issues.append(f"font@{i}: {info['font']} | {t[:60]}")
                if info["size"] and info["size"] not in ("177800", "203200", "None"):  # 14pt=177800, 16pt=203200
                    issues.append(f"size@{i}: {info['size']} | {t[:60]}")
            lines.append(f"Issues count: {len(issues)}")
            for x in issues[:80]:
                lines.append(x)

    with open(OUT_PATH, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print(f"Written {OUT_PATH}")


if __name__ == "__main__":
    main()
