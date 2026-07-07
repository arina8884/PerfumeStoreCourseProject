# -*- coding: utf-8 -*-
"""Исправление порядка содержания (forward insert)."""
import re
import shutil
from pathlib import Path
from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.shared import Cm, Pt

DOC = Path(r"c:\Users\user\Desktop\PerfumeStore_NEW\PerfumeStoreCourseProject\Documentation\ПЗКурсовая_ГОТОВО.docx")
OUT = Path(r"c:\Users\user\Desktop\PerfumeStore_NEW\PerfumeStoreCourseProject\Documentation\ПЗКурсовая.docx")

APPENDIX_RE = re.compile(r"^ПРИЛОЖЕНИЕ\s+[А-ЯA-Z]$")

def is_heading(p, t):
    return p.style.name == "Заголовок ПЗ 1" and bool(t)

def pages_map(doc):
    page = 1
    d = {}
    for i, p in enumerate(doc.paragraphs):
        if i > 0 and "lastRenderedPageBreak" in p._element.xml:
            page += 1
        d[i] = page
    return d

def collect(doc, pages, start):
    out = []
    for i, p in enumerate(doc.paragraphs):
        if i < start: continue
        t = p.text.strip()
        if not t or "\t" in t: continue
        if is_heading(p, t) and t.upper() != "СОДЕРЖАНИЕ":
            lvl = 2 if re.match(r"^\d+\.\d+", t) else (1 if re.match(r"^\d+\s", t) else 0)
            if not re.match(r"^\d", t) and t.lower() not in ("введение","заключение"): lvl = 2
            out.append((t, pages[i], lvl))
        if t.lower() == "список использованных источников":
            out.append((t, pages[i], 0))
        if APPENDIX_RE.match(t):
            title = ""
            for j in range(i+1, min(i+6, len(doc.paragraphs))):
                t2 = doc.paragraphs[j].text.strip()
                if t2 and t2 not in ("(обязательное)","(рекомендуемое)","(справочное)"):
                    title = t2; break
            out.append((f"{t} {title}".strip(), pages[i], 0))
    return out

doc = Document(DOC)
ci = ii = None
for i,p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t.upper()=="СОДЕРЖАНИЕ": ci=i
    if ii is None and ci is not None and i>ci and is_heading(p,t) and t.lower().startswith("введение"): ii=i

body = doc.element.body
ce, ie = doc.paragraphs[ci]._element, doc.paragraphs[ii]._element
rem=False
for ch in list(body):
    if ch is ce: rem=True; continue
    if ch is ie: break
    if rem: body.remove(ch)

pages = pages_map(doc)
ii2 = next(i for i,p in enumerate(doc.paragraphs) if is_heading(p,p.text.strip()) and p.text.strip().lower().startswith("введение"))
entries = collect(doc, pages, ii2)
intro = doc.paragraphs[ii2]
tab = Cm(16.5)
for title, page, lvl in entries:  # forward!
    np = intro.insert_paragraph_before(f"{'  '*lvl}{title}\t{page}")
    np.style = doc.styles['Normal']
    np.paragraph_format.tab_stops.add_tab_stop(tab, WD_TAB_ALIGNMENT.RIGHT)
    np.paragraph_format.space_after = Pt(0)

doc.save(DOC)
shutil.copy2(DOC, OUT)
print('done')
for i in range(50, 76):
    print(i, doc.paragraphs[i].text)
