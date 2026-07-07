# -*- coding: utf-8 -*-
from docx import Document
from collections import Counter

path = r"c:\Users\user\Desktop\PerfumeStore_NEW\PerfumeStoreCourseProject\Documentation\ПЗКурсовая.docx"
doc = Document(path)
styles = Counter(p.style.name for p in doc.paragraphs)
print(styles)
for i,p in enumerate(doc.paragraphs):
    if 'toc' in p.style.name.lower() or 'TOC' in p.style.name:
        print(i, p.style.name, repr(p.text[:80]))
